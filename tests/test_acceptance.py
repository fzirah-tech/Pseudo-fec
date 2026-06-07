"""Tests d'acceptation (cf. §12 du brief)."""

import os
import shutil

import pytest

from pseudo_fec.mapping import TOKEN_RE, Mapping
from pseudo_fec.pseudonymize import Pseudonymizer, read_fec, write_fec
from pseudo_fec.reidentify import reidentify_file
from pseudo_fec.verify import verify

DATA = os.path.join(os.path.dirname(__file__), "data", "FEC_DEMO_2024.txt")


def _sum_debit_credit(fec):
    f_d = fec.field("Debit")
    f_c = fec.field("Credit")

    def val(s):
        s = str(s or "").strip().replace(" ", "").replace(" ", "")
        s = s.replace(",", ".")
        return float(s) if s else 0.0

    debit = sum(val(r.get(f_d, "")) for r in fec.rows)
    credit = sum(val(r.get(f_c, "")) for r in fec.rows)
    return round(debit, 2), round(credit, 2)


@pytest.fixture
def pseudonymized(tmp_path):
    out = str(tmp_path / "FEC_pseudo.txt")
    fec = read_fec(DATA)
    mapping = Mapping()
    fec_out, mapping = Pseudonymizer().run(fec, mapping)
    write_fec(fec_out, out)
    return out, mapping


# 1. Round-trip docx / xlsx
def _token_for(mapping, real):
    for e in mapping.entries.values():
        if e.valeur_reelle == real:
            return e.pseudo
    raise AssertionError(f"Pas de pseudo pour {real}")


def test_roundtrip_docx(pseudonymized, tmp_path):
    from docx import Document

    _, mapping = pseudonymized
    token = _token_for(mapping, "DUPONT SARL")
    original = "Le chiffre d'affaires de DUPONT SARL progresse en 2024."
    pseudo_text = original.replace("DUPONT SARL", token)

    src = str(tmp_path / "analyse.docx")
    doc = Document()
    doc.add_paragraph(pseudo_text)
    doc.save(src)

    out = str(tmp_path / "analyse_reid.docx")
    reidentify_file(src, mapping, out)

    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert text == original


def test_roundtrip_xlsx(pseudonymized, tmp_path):
    from openpyxl import Workbook, load_workbook

    _, mapping = pseudonymized
    token = _token_for(mapping, "MARTIN ET FILS")
    original = "Total achats MARTIN ET FILS"
    pseudo_text = original.replace("MARTIN ET FILS", token)

    src = str(tmp_path / "tableau.xlsx")
    wb = Workbook()
    wb.active["A1"] = pseudo_text
    wb.active["B1"] = 1234.56  # un nombre ne doit pas être altéré
    wb.save(src)

    out = str(tmp_path / "tableau_reid.xlsx")
    reidentify_file(src, mapping, out)

    ws = load_workbook(out).active
    assert ws["A1"].value == original
    assert ws["B1"].value == 1234.56


# 2. Cohérence : un même tiers = un même pseudo partout
def test_consistency(pseudonymized):
    out, _ = pseudonymized
    fec = read_fec(out)
    f_auxlib = fec.field("CompAuxLib")
    tokens = {
        r[f_auxlib]
        for r in fec.rows
        if str(r.get(f_auxlib, "")).strip() and r.get("JournalCode") in ("VT", "BQ")
        and "CLI" in str(r.get(f_auxlib, ""))
    }
    # DUPONT SARL (aux CDUPONT) apparaît dans 2 écritures -> 1 seul pseudo.
    assert len(tokens) == 1


# 3. Préservation des montants, dates et comptes
def test_preservation(pseudonymized):
    out, _ = pseudonymized
    src_fec = read_fec(DATA)
    out_fec = read_fec(out)
    assert _sum_debit_credit(src_fec) == _sum_debit_credit(out_fec)

    # Dates et comptes non-tiers inchangés.
    f_date = src_fec.field("EcritureDate")
    f_cpt = src_fec.field("CompteNum")
    for s, o in zip(src_fec.rows, out_fec.rows):
        assert s[f_date] == o[f_date]
        # Les comptes généraux (non tiers) sont strictement préservés.
        if s[f_cpt][:3] not in ("401", "411", "421"):
            assert s[f_cpt] == o[f_cpt]


# 4. Anti-fuite
def test_anti_leak(pseudonymized):
    out, mapping = pseudonymized
    leaks, _ = verify(out, mapping)
    assert leaks == [], f"Fuites détectées : {leaks}"


# 5. CompteNum concaténé : 411DUPON -> 411CLI000x, plus aucune occurrence de DUPON
def test_comptenum(pseudonymized):
    out, _ = pseudonymized
    with open(out, encoding="utf-8") as f:
        content = f.read()
    assert "DUPON" not in content
    assert "DURAND" not in content
    fec = read_fec(out)
    f_cpt = fec.field("CompteNum")
    comptes = [r[f_cpt] for r in fec.rows]
    assert any(TOKEN_RE.search("§" + c.replace("411", "", 1) + "§") or c.startswith("411CLI")
               for c in comptes)
    assert any(c.startswith("411CLI") for c in comptes)
    assert any(c.startswith("421SAL") for c in comptes)


# 6. Idempotence : double passage = résultat identique
def test_idempotence(pseudonymized, tmp_path):
    out1, mapping = pseudonymized
    n_entities = len(mapping.entries)

    fec1 = read_fec(out1)
    fec2, mapping = Pseudonymizer().run(fec1, mapping)
    out2 = str(tmp_path / "FEC_pseudo2.txt")
    write_fec(fec2, out2)

    with open(out1, encoding="utf-8") as f:
        c1 = f.read()
    with open(out2, encoding="utf-8") as f:
        c2 = f.read()
    assert c1 == c2
    assert len(mapping.entries) == n_entities


# Persistance incrémentale du mapping
def test_incremental_mapping(tmp_path):
    map_path = str(tmp_path / "mapping.csv")
    fec = read_fec(DATA)
    m1 = Mapping()
    _, m1 = Pseudonymizer().run(fec, m1)
    m1.save(map_path)
    tokens_1 = {e.valeur_reelle: e.pseudo for e in m1.entries.values()}

    # Re-run en chargeant le mapping existant : pseudos réutilisés.
    fec2 = read_fec(DATA)
    m2 = Mapping.load(map_path)
    _, m2 = Pseudonymizer().run(fec2, m2)
    tokens_2 = {e.valeur_reelle: e.pseudo for e in m2.entries.values()}

    assert tokens_1 == tokens_2
