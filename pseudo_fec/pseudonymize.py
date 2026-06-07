"""Pseudonymisation d'un FEC : I/O, détection des tiers, réécriture.

Conception (cf. brief) :
- on pseudonymise le QUI (noms, codes auxiliaires identifiants) ;
- on préserve le QUOI/COMBIEN/QUAND (montants, dates, comptes PCG, journaux,
  lettrage) — aucune de ces colonnes n'est reformatée.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from .mapping import (
    BARE_COMPTE_RE,
    Mapping,
    is_pseudo_value,
    normalize,
)
from .replacer import Replacer

# 18 champs normalisés DGFiP, séparateur pipe.
FEC_FIELDS = [
    "JournalCode", "JournalLib", "EcritureNum", "EcritureDate",
    "CompteNum", "CompteLib", "CompAuxNum", "CompAuxLib",
    "PieceRef", "PieceDate", "EcritureLib", "Debit", "Credit",
    "EcritureLet", "DateLet", "ValidDate", "Montantdevise", "Idevise",
]

CANDIDATE_DELIMITERS = ["|", "\t", ";"]
# Sans BOM, on privilégie utf-8 (réécriture sans BOM) ; utf-8-sig n'est retenu
# que si un BOM est réellement présent, afin de préserver l'encodage d'origine.
CANDIDATE_ENCODINGS = ["utf-8", "cp1252", "iso-8859-1"]

# Préfixe de compte (3 car.) -> type de tiers.
ACCOUNT_TYPE = {
    "401": "FRN", "403": "FRN", "404": "FRN", "408": "FRN", "409": "FRN",
    "411": "CLI", "413": "CLI", "416": "CLI", "418": "CLI", "419": "CLI",
    "421": "SAL", "422": "SAL", "425": "SAL", "427": "SAL", "428": "SAL",
}
DEFAULT_TIER_TYPE = "TIE"
# Préfixes considérés comme comptes de tiers (sous-comptes nominatifs).
TIER_ACCOUNT_PREFIXES = tuple(ACCOUNT_TYPE.keys()) + ("455", "467", "108")


# ---------------------------------------------------------------------- #
# I/O FEC
# ---------------------------------------------------------------------- #
class FecData:
    def __init__(self, fieldnames, rows, delimiter, encoding, fmt, raw_text=""):
        self.fieldnames = fieldnames
        self.rows = rows            # list[dict] ; valeurs natives (str/num/date)
        self.delimiter = delimiter
        self.encoding = encoding
        self.fmt = fmt              # "text" ou "xlsx"
        self.raw_text = raw_text    # contenu brut (texte) pour l'anti-collision

    def field(self, logical: str) -> Optional[str]:
        target = logical.lower().replace("_", "").replace(" ", "")
        for name in self.fieldnames:
            if name and name.lower().replace("_", "").replace(" ", "") == target:
                return name
        return None


def _detect_encoding(path: str) -> str:
    import codecs

    with open(path, "rb") as f:
        head = f.read(3)
    if head.startswith(codecs.BOM_UTF8):
        return "utf-8-sig"
    for enc in CANDIDATE_ENCODINGS:
        try:
            with open(path, encoding=enc) as f:
                f.read()
            return enc
        except (UnicodeDecodeError, UnicodeError):
            continue
    return "iso-8859-1"


def _detect_delimiter(header: str) -> str:
    counts = {d: header.count(d) for d in CANDIDATE_DELIMITERS}
    best = max(counts, key=counts.get)
    return best if counts[best] else "|"


def read_fec(path: str) -> FecData:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        return _read_fec_xlsx(path)
    return _read_fec_text(path)


def _read_fec_text(path: str) -> FecData:
    import csv

    encoding = _detect_encoding(path)
    with open(path, encoding=encoding, newline="") as f:
        raw_text = f.read()
    delimiter = _detect_delimiter(raw_text.splitlines()[0] if raw_text else "")
    with open(path, encoding=encoding, newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        fieldnames = reader.fieldnames or []
        rows = [
            {k: (v if v is not None else "") for k, v in r.items()}
            for r in reader
        ]
    return FecData(fieldnames, rows, delimiter, encoding, "text", raw_text)


def _read_fec_xlsx(path: str) -> FecData:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    ws = wb["FEC"] if "FEC" in wb.sheetnames else wb.active
    it = ws.iter_rows(values_only=True)
    header = next(it, None)
    if header is None:
        raise ValueError("Feuille FEC vide.")
    fieldnames = [str(h) if h is not None else "" for h in header]
    rows = []
    raw_parts = ["|".join(fieldnames)]
    for values in it:
        row = {}
        cells = []
        for i, name in enumerate(fieldnames):
            v = values[i] if i < len(values) else None
            v = "" if v is None else v
            row[name] = v
            cells.append("" if v == "" else str(v))
        if any(c != "" for c in cells):
            rows.append(row)
            raw_parts.append("|".join(cells))
    return FecData(fieldnames, rows, "|", "utf-8", "xlsx", "\n".join(raw_parts))


def write_fec(fec: FecData, path: str) -> None:
    if fec.fmt == "xlsx":
        _write_fec_xlsx(fec, path)
    else:
        _write_fec_text(fec, path)


def _write_fec_text(fec: FecData, path: str) -> None:
    import csv

    with open(path, "w", encoding=fec.encoding, newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=fec.fieldnames,
            delimiter=fec.delimiter,
            restval="",
            extrasaction="ignore",
            lineterminator="\r\n",
        )
        writer.writeheader()
        for row in fec.rows:
            writer.writerow(row)


def _write_fec_xlsx(fec: FecData, path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "FEC"
    ws.append(fec.fieldnames)
    for row in fec.rows:
        ws.append([row.get(name, "") for name in fec.fieldnames])
    wb.save(path)


# ---------------------------------------------------------------------- #
# Détection / réécriture
# ---------------------------------------------------------------------- #
def _account_type(compte: str) -> str:
    compte = (compte or "").strip()
    return ACCOUNT_TYPE.get(compte[:3], DEFAULT_TIER_TYPE)


def _is_subaccount(compte: str) -> bool:
    c = (compte or "").strip()
    if len(c) <= 3 or not c.startswith(TIER_ACCOUNT_PREFIXES):
        return False
    return c[3:].strip("0") != ""


def _log(message: str) -> None:
    """Journalisation locale, sans valeur sensible."""
    try:
        os.makedirs("logs", exist_ok=True)
        with open(os.path.join("logs", "pseudo_fec.log"), "a", encoding="utf-8") as f:
            f.write(f"{datetime.now().isoformat()} {message}\n")
    except OSError:
        pass


class Pseudonymizer:
    def __init__(self, include_pieceref: bool = False, ascii_mode: bool = False):
        self.include_pieceref = include_pieceref
        self.ascii_mode = ascii_mode

    def _row_entity(self, fec: FecData, row: dict):
        """Retourne (kind, key, type, valeur_reelle, code_aux) ou None.

        kind ∈ {"aux", "compte"}.
        """
        f_auxnum = fec.field("CompAuxNum")
        f_auxlib = fec.field("CompAuxLib")
        f_cptnum = fec.field("CompteNum")
        f_cptlib = fec.field("CompteLib")

        auxnum = str(row.get(f_auxnum, "") if f_auxnum else "").strip()
        auxlib = str(row.get(f_auxlib, "") if f_auxlib else "").strip()
        cptnum = str(row.get(f_cptnum, "") if f_cptnum else "").strip()
        cptlib = str(row.get(f_cptlib, "") if f_cptlib else "").strip()

        if auxnum or auxlib:
            # Lignes déjà pseudonymisées : on ne recrée pas d'entité.
            if is_pseudo_value(auxnum) or is_pseudo_value(auxlib):
                return None
            type_ = _account_type(cptnum)
            key = ("aux", auxnum.upper() or normalize(auxlib))
            return ("aux", key, type_, auxlib, auxnum)

        if _is_subaccount(cptnum):
            if is_pseudo_value(cptnum) or BARE_COMPTE_RE.fullmatch(cptnum):
                return None
            type_ = _account_type(cptnum)
            key = ("cpt", cptnum.upper())
            return ("compte", key, type_, cptlib, cptnum)

        return None

    def run(
        self,
        fec: FecData,
        mapping: Mapping,
    ) -> Tuple[FecData, Mapping]:
        mapping.ascii_mode = self.ascii_mode

        f_auxnum = fec.field("CompAuxNum")
        f_auxlib = fec.field("CompAuxLib")
        f_cptnum = fec.field("CompteNum")
        f_cptlib = fec.field("CompteLib")
        f_ecrlib = fec.field("EcritureLib")
        f_pieceref = fec.field("PieceRef")

        # Passe 1 : détection et attribution des pseudonymes.
        new_count = 0
        for row in fec.rows:
            ent = self._row_entity(fec, row)
            if ent is None:
                continue
            kind, key, type_, valeur, code_aux = ent
            before = len(mapping.entries)
            champ = "CompteLib" if kind == "compte" else "CompAuxLib"
            mapping.get_or_create(
                key, type_, valeur, champ, code_aux, fec.raw_text
            )
            if len(mapping.entries) > before:
                new_count += 1

        # Masquage complet optionnel des PieceRef.
        ref_tokens: Dict[str, str] = {}
        if self.include_pieceref and f_pieceref:
            for row in fec.rows:
                val = str(row.get(f_pieceref, "") or "").strip()
                if not val or is_pseudo_value(val):
                    continue
                key = ("ref", normalize(val))
                entry = mapping.get_or_create(
                    key, "REF", val, "PieceRef", "", fec.raw_text
                )
                ref_tokens[val] = entry.pseudo

        # Réplacement en texte libre (libellés/codes de tiers connus).
        replacer = Replacer(mapping.forward_pairs())

        # Passe 2 : réécriture.
        freetext_fields = [f for f in (f_ecrlib, f_pieceref, f_cptlib) if f]
        for row in fec.rows:
            ent = self._row_entity(fec, row)
            if ent is not None:
                kind, key, _, _, _ = ent
                entry = mapping.entries[key]
                entry.occurrences += 1
                if kind == "aux":
                    if f_auxlib and str(row.get(f_auxlib, "")).strip():
                        row[f_auxlib] = entry.pseudo
                    if f_auxnum and str(row.get(f_auxnum, "")).strip():
                        row[f_auxnum] = entry.pseudo
                else:  # compte : on conserve la racine collective.
                    cptnum = str(row.get(f_cptnum, "")).strip()
                    root = cptnum[:3]
                    if f_cptnum:
                        row[f_cptnum] = f"{root}{entry.bare}"
                    if f_cptlib:
                        row[f_cptlib] = entry.pseudo

            # Texte libre.
            if replacer:
                for fld in freetext_fields:
                    val = row.get(fld, "")
                    if isinstance(val, str) and val:
                        row[fld] = replacer.apply(val)

            # Masquage complet PieceRef.
            if ref_tokens and f_pieceref:
                val = str(row.get(f_pieceref, "") or "").strip()
                if val in ref_tokens:
                    row[f_pieceref] = ref_tokens[val]

        mapping.meta.update(
            {
                "updated": datetime.now().isoformat(timespec="seconds"),
                "entities_total": len(mapping.entries),
            }
        )
        _log(
            f"pseudonymize: rows={len(fec.rows)} entities_total={len(mapping.entries)} "
            f"new_entities={new_count} include_pieceref={self.include_pieceref} "
            f"ascii={self.ascii_mode}"
        )
        return fec, mapping
