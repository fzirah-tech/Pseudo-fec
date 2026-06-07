"""Ré-identification locale : remplace les pseudonymes par les valeurs réelles
dans les livrables (.docx, .xlsx, .txt/.md/.csv).

Le remplacement porte sur le token complet délimité (ex. §CLI0001§), ce qui
élimine toute collision partielle (CLI0001 vs CLI00010).
"""

from __future__ import annotations

import os
from typing import List, Tuple

from .mapping import Mapping
from .replacer import Replacer


def _replacer(mapping: Mapping) -> Replacer:
    return Replacer(mapping.reverse_pairs(), min_length=2)


# ---------------------------------------------------------------------- #
# Word (.docx) — fusion des runs adjacents de mêmes propriétés
# ---------------------------------------------------------------------- #
def _run_signature(run):
    font = run.font
    color = None
    try:
        if font.color and font.color.type is not None:
            color = str(font.color.rgb)
    except Exception:
        color = None
    return (
        run.bold,
        run.italic,
        run.underline,
        font.name,
        font.size,
        color,
        run.style.name if run.style is not None else None,
    )


def _merge_adjacent_runs(paragraph) -> None:
    """Fusionne les runs adjacents partageant les mêmes propriétés.

    Un pseudonyme peut être scindé entre plusieurs runs : la fusion permet de
    retrouver le token complet avant remplacement, tout en préservant la mise
    en forme là où les runs voisins partagent le même style.
    """
    runs = paragraph.runs
    if len(runs) < 2:
        return
    i = 0
    while i < len(runs) - 1:
        cur, nxt = runs[i], runs[i + 1]
        if _run_signature(cur) == _run_signature(nxt):
            cur.text = (cur.text or "") + (nxt.text or "")
            nxt.text = ""
            # Retire l'élément vide du XML pour garder le document propre.
            nxt._element.getparent().remove(nxt._element)
            runs = paragraph.runs
        else:
            i += 1


def _replace_in_paragraph(paragraph, replacer: Replacer) -> int:
    n = 0
    _merge_adjacent_runs(paragraph)
    for run in paragraph.runs:
        if run.text and replacer.contains(run.text):
            run.text = replacer.apply(run.text)
            n += 1
    # Filet de sécurité : token réparti sur des runs de styles différents.
    if replacer.contains(paragraph.text):
        new_text = replacer.apply(paragraph.text)
        if paragraph.runs:
            paragraph.runs[0].text = new_text
            for run in paragraph.runs[1:]:
                run.text = ""
        n += 1
    return n


def _process_docx(path: str, out_path: str, replacer: Replacer) -> int:
    from docx import Document

    doc = Document(path)
    total = 0

    def do_paragraphs(paragraphs):
        nonlocal total
        for p in paragraphs:
            total += _replace_in_paragraph(p, replacer)

    def do_tables(tables):
        for table in tables:
            for trow in table.rows:
                for cell in trow.cells:
                    do_paragraphs(cell.paragraphs)
                    do_tables(cell.tables)

    do_paragraphs(doc.paragraphs)
    do_tables(doc.tables)
    for section in doc.sections:
        do_paragraphs(section.header.paragraphs)
        do_paragraphs(section.footer.paragraphs)
        do_tables(section.header.tables)
        do_tables(section.footer.tables)

    doc.save(out_path)
    return total


# ---------------------------------------------------------------------- #
# Excel (.xlsx)
# ---------------------------------------------------------------------- #
def _process_xlsx(path: str, out_path: str, replacer: Replacer) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                # Ne pas altérer les formules (chaînes commençant par '=').
                if isinstance(val, str) and not val.startswith("="):
                    if replacer.contains(val):
                        cell.value = replacer.apply(val)
                        n += 1
    wb.save(out_path)
    return n


# ---------------------------------------------------------------------- #
# Texte / Markdown / CSV
# ---------------------------------------------------------------------- #
def _process_text(path: str, out_path: str, replacer: Replacer) -> int:
    import codecs

    with open(path, "rb") as fb:
        if fb.read(3).startswith(codecs.BOM_UTF8):
            with open(path, encoding="utf-8-sig") as f:
                content = f.read()
            new_content = replacer.apply(content)
            with open(out_path, "w", encoding="utf-8-sig") as f:
                f.write(new_content)
            return 0 if new_content == content else 1

    encodings = ["utf-8", "cp1252", "iso-8859-1"]
    content, used = None, "utf-8"
    for enc in encodings:
        try:
            with open(path, encoding=enc) as f:
                content = f.read()
            used = enc
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if content is None:
        with open(path, encoding="iso-8859-1") as f:
            content = f.read()
        used = "iso-8859-1"
    new_content = replacer.apply(content)
    with open(out_path, "w", encoding=used) as f:
        f.write(new_content)
    return 0 if new_content == content else 1


# ---------------------------------------------------------------------- #
# Point d'entrée
# ---------------------------------------------------------------------- #
def reidentify_file(path: str, mapping: Mapping, out_path: str) -> Tuple[str, int]:
    replacer = _replacer(mapping)
    ext = os.path.splitext(path)[1].lower()
    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    if ext == ".docx":
        n = _process_docx(path, out_path, replacer)
    elif ext == ".xlsx":
        n = _process_xlsx(path, out_path, replacer)
    elif ext in (".txt", ".md", ".csv", ".tsv"):
        n = _process_text(path, out_path, replacer)
    else:
        raise ValueError(
            f"Format non pris en charge : {ext} "
            f"(.docx, .xlsx, .txt, .md, .csv)"
        )
    return out_path, n
